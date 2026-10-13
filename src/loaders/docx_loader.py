from docx import Document as DocxDocument
from models.document import Document


def extract_text(file_path):

    doc = DocxDocument(file_path)

    paragraphs = []

    for paragraph in doc.paragraphs:

        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    text = "\n".join(paragraphs)

    return Document(
        source_file=file_path,
        source_type="docx",
        content=text
    )