from openpyxl import load_workbook
from models.document import Document


def extract_text(file_path):

    workbook = load_workbook(
        file_path,
        read_only=True,
        data_only=True
    )

    sheets_text = []

    try:

        for sheet in workbook.worksheets:

            rows = []

            for row in sheet.iter_rows(values_only=True):

                cells = [
                    str(cell).strip()
                    for cell in row
                    if cell is not None and str(cell).strip()
                ]

                if cells:
                    rows.append(" | ".join(cells))

            if rows:
                sheets_text.append(
                    f"[Hoja: {sheet.title}]\n" + "\n".join(rows)
                )

    finally:

        workbook.close()

    content = "\n\n".join(sheets_text)

    return Document(
        source_file=file_path,
        source_type="xlsx",
        content=content
    )
