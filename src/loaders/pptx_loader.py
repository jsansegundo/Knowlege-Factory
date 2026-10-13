from pptx import Presentation
from models.document import Document


def extract_text(file_path):

    presentation = Presentation(file_path)

    slides_text = []

    for slide in presentation.slides:

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                text = shape.text.strip()

                if text:
                    slides_text.append(text)

    content = "\n".join(slides_text)

    return Document(
        source_file=file_path,
        source_type="pptx",
        content=content
    )