"""
Genera las sesiones de un workshop a partir de la base de conocimiento (RAG).

Todo el contenido sale de consultas al RAG: el bloque Zero Trust común y los
dominios elegidos para el cliente. El modelo solo organiza lo recuperado en
diapositivas y notas de orador, citando las fuentes.

Uso:
    python src/build_workshop_from_rag.py <perfil.json> [carpeta_salida]
"""

import json
import os
import sys
from pathlib import Path

import faiss
from dotenv import load_dotenv
from pptx import Presentation
from pptx.util import Pt

from knowledge_chat import enrich_chunks_with_registry_metadata
from knowledge_pipeline import retrieve
from openai_client import create_openai_client
from registry import load_registry

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_OUTPUT = PROJECT_ROOT / "outputs" / "frontier_workshop"

ZERO_TRUST_QUERIES = [
    "Zero Trust Workshop pillars and progress report",
    "Zero Trust Workshop delivery steps and assessment",
]

SESSION_TOPICS = {
    "01_Kickoff.pptx": "kickoff, objectives, scope, customer prerequisites and discovery of a security engagement",
    "02_Configuracion_y_Zero_Trust.pptx": "configuration and delivery steps of the security modules",
    "03_Recogida_de_Evidencias_e_Informe.pptx": "evidence collection, findings, results presentation and report for a security engagement",
}

SLIDE_PROMPT = """
Eres formador de un workshop de seguridad. Con la documentación recuperada,
escribe las diapositivas de la sesión indicada en español.

Devuelve SOLO un JSON con esta forma:
{{"slides": [{{"title": "...", "bullets": ["..."], "notes": "...", "sources": ["documento"]}}]}}

Reglas:
- Usa solo la información del CONTEXTO. Si algo no aparece, no lo inventes.
- Entre 4 y 6 diapositivas. Cada una con 3 a 5 viñetas.
- Las notas de orador explican cómo presentar la diapositiva.
- En "sources" pon los nombres de los documentos usados.

SESIÓN: {session}
TEMA: {topic}
DOMINIO DEL CLIENTE: {domain}

CONTEXTO:
{context}
"""


def load_knowledge(client):
    index = faiss.read_index(str(PROJECT_ROOT / "indexes" / "knowledge.index"))
    with open(PROJECT_ROOT / "manifests" / "all_chunks.json", "r", encoding="utf-8") as f:
        chunks = json.load(f)
    enrich_chunks_with_registry_metadata(chunks, load_registry())
    return index, chunks


def retrieve_for(client, index, chunks, queries):
    seen = []
    selected = []
    for query in queries:
        for chunk in retrieve(client, index, chunks, query):
            key = (chunk["document_name"], chunk["chunk_number"])
            if key not in seen:
                seen.append(key)
                selected.append(chunk)
    return selected


def format_context(chunks):
    return "\n\n".join(
        f"[Documento: {c['document_name']} | Fragmento: {c['chunk_number']}]\n{c['content']}"
        for c in chunks
    )


def generate_slides(client, session, topic, domain, chunks):
    prompt = SLIDE_PROMPT.format(
        session=session, topic=topic, domain=domain, context=format_context(chunks)
    )
    raw = client.responses.create(
        model=os.getenv("AZURE_OPENAI_DEPLOYMENT"),
        input=prompt,
    ).output_text.strip()
    start, end = raw.find("{"), raw.rfind("}")
    return json.loads(raw[start:end + 1])["slides"]


def add_slide(presentation, slide_data):
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = slide_data["title"]
    body = slide.placeholders[1].text_frame
    body.clear()
    for position, bullet in enumerate(slide_data["bullets"]):
        paragraph = body.paragraphs[0] if position == 0 else body.add_paragraph()
        paragraph.text = bullet
        paragraph.font.size = Pt(20)
    sources = ", ".join(slide_data.get("sources", []))
    notes = slide_data["notes"]
    if sources:
        notes += f"\n\nFuentes: {sources}"
    slide.notes_slide.notes_text_frame.text = notes


def build_deck(client_name, session_file, slides, output_folder):
    presentation = Presentation()
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = f"{session_file.replace('.pptx', '').replace('_', ' ')} · {client_name}"
    title_slide.placeholders[1].text = "Frontier Accelerate for Security · Envisioning & POC"
    title_slide.notes_slide.notes_text_frame.text = "Presenta la sesión y su objetivo."

    for slide_data in slides:
        add_slide(presentation, slide_data)

    output_folder.mkdir(parents=True, exist_ok=True)
    path = output_folder / session_file
    presentation.save(path)
    return path


def main():
    if len(sys.argv) < 2:
        raise SystemExit("Uso: python src/build_workshop_from_rag.py <perfil.json> [carpeta_salida]")

    load_dotenv(PROJECT_ROOT / ".env")
    profile = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    output_folder = Path(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_OUTPUT

    client = create_openai_client()
    index, chunks = load_knowledge(client)

    client_name = profile["client"]
    domain = profile["domain"]
    domain_queries = profile["domain_queries"]

    zero_trust_chunks = retrieve_for(client, index, chunks, ZERO_TRUST_QUERIES)
    domain_chunks = retrieve_for(client, index, chunks, domain_queries)

    for session_file, topic in SESSION_TOPICS.items():
        if session_file.startswith("02"):
            session_chunks = zero_trust_chunks + domain_chunks
        else:
            session_chunks = domain_chunks
        slides = generate_slides(client, session_file, topic, domain, session_chunks)
        path = build_deck(client_name, session_file, slides, output_folder / client_name.lower())
        print(f"Creada: {path} ({len(slides)} diapositivas)")


if __name__ == "__main__":
    main()
